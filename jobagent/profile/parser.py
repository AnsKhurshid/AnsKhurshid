"""Read a resume (DOCX or PDF) into ordered lines with light structure hints."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from jobagent.utils.text import clean_text

BULLET_RE = re.compile(r"^\s*([•●▪◦‣∙·\-–—*➢►✓■□o])\s+")


class ResumeParseError(Exception):
    pass


@dataclass
class Line:
    text: str
    bullet: bool = False
    heading_hint: bool = False   # styled as heading, or short all-caps / bold line


def _docx_lines(path: Path) -> list[Line]:
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document(str(path))
    lines: list[Line] = []

    def para_line(p) -> Line | None:
        text = clean_text(p.text)
        if not text:
            return None
        style = (p.style.name or "").lower() if p.style is not None else ""
        is_list = "list" in style or p._p.pPr is not None and p._p.pPr.find(qn("w:numPr")) is not None
        bold = bool(p.runs) and all((r.bold or (r.style is not None and r.style.font.bold)) for r in p.runs if r.text.strip())
        heading = style.startswith("heading") or style == "title" or (bold and len(text) < 40) or (
            text.isupper() and len(text) < 40)
        m = BULLET_RE.match(text)
        if m:
            text, is_list = text[m.end():].strip(), True
        return Line(text=text, bullet=is_list and not style.startswith("heading"), heading_hint=heading and not is_list)

    # Walk body in document order so tables (common in resume templates) are not lost.
    body = doc.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            if (ln := para_line(Paragraph(child, doc))):
                lines.append(ln)
        elif child.tag == qn("w:tbl"):
            for row in Table(child, doc).rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        if (ln := para_line(p)):
                            lines.append(ln)
    # Headers often carry name/contact details.
    for section in doc.sections:
        for p in section.header.paragraphs:
            if (ln := para_line(p)):
                lines.insert(0, ln)
    return lines


def _pdf_lines(path: Path) -> list[Line]:
    import pdfplumber

    lines: list[Line] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=1.5, y_tolerance=3) or ""
            for raw in clean_text(text).splitlines():
                if not raw.strip():
                    continue
                m = BULLET_RE.match(raw)
                if m:
                    lines.append(Line(text=raw[m.end():].strip(), bullet=True))
                else:
                    t = raw.strip()
                    lines.append(Line(text=t, heading_hint=t.isupper() and len(t) < 40))
    return merge_wrapped_bullets(lines)


def merge_wrapped_bullets(lines: list[Line]) -> list[Line]:
    """Join bullet continuation lines (PDFs wrap long bullets onto new lines).

    A wrapped line may start with a capital ("...multiple MLS / API sources with..."), so the main
    signal is that the previous bullet has not finished its sentence. Lines that look like a section
    heading or a dated role header are never joined.
    """
    from jobagent.profile.extractor import DATE_RANGE_RE, _section_of
    merged: list[Line] = []
    for ln in lines:
        prev = merged[-1] if merged else None
        continues = (prev is not None and prev.bullet and not ln.bullet and not ln.heading_hint
                     and _section_of(ln) is None and not DATE_RANGE_RE.search(ln.text)
                     and (ln.text[:1].islower() or not prev.text.rstrip().endswith((".", "!", "?", ":"))))
        if continues:
            prev.text += " " + ln.text
        else:
            merged.append(ln)
    return merged


def read_resume(path: Path) -> list[Line]:
    path = Path(path)
    if not path.exists():
        raise ResumeParseError(f"resume not found: {path}")
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            lines = _docx_lines(path)
        elif suffix == ".pdf":
            lines = _pdf_lines(path)
        elif suffix in (".txt", ".md"):
            lines = []
            for raw in clean_text(path.read_text(encoding="utf-8")).splitlines():
                if raw.strip():
                    m = BULLET_RE.match(raw)
                    lines.append(Line(text=raw[m.end():].strip(), bullet=True) if m else Line(text=raw.strip()))
        else:
            raise ResumeParseError(f"unsupported resume format '{suffix}' (use .docx or .pdf)")
    except ResumeParseError:
        raise
    except Exception as exc:  # corrupted / encrypted files
        raise ResumeParseError(f"could not read {path.name}: {exc}") from exc
    if sum(len(l.text) for l in lines) < 200:
        raise ResumeParseError(
            f"very little text extracted from {path.name}; if it is a scanned PDF, please provide a DOCX "
            "or a text-based PDF (OCR is not supported)")
    return lines
