"""Render a TailoredResume as a clean, single-column, ATS-friendly DOCX."""
from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from jobagent.resume.model import TailoredResume

FONT = "Calibri"


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "808080")):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def _heading(doc, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    _bottom_border(p)


def _bullet(doc, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.add_run(text)


def render_docx(resume: TailoredResume, title: str) -> bytes:
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.6)
        s.left_margin = s.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(resume.name)
    r.bold = True
    r.font.size = Pt(18)
    if resume.headline:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(resume.headline).font.size = Pt(11.5)
    c = resume.contact
    contact = " | ".join(x for x in (c.location, c.email, c.phone, c.linkedin, c.github, c.website) if x)
    if contact:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(contact).font.size = Pt(9.5)

    _heading(doc, "Professional Summary")
    doc.add_paragraph(resume.summary)

    _heading(doc, "Technical Skills")
    for g in resume.skill_groups:
        p = doc.add_paragraph()
        p.add_run(f"{g.label}: ").bold = True
        p.add_run(", ".join(g.skills))
    if resume.exploring:
        p = doc.add_paragraph()
        p.add_run("Currently learning: ").bold = True
        p.add_run(", ".join(resume.exploring))

    _heading(doc, "Professional Experience")
    for role in resume.experience:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.1), alignment=2)  # right-aligned dates
        p.add_run(role.title).bold = True
        p.add_run(f" | {role.employer}" + (f", {role.location}" if role.location else ""))
        dates = " – ".join(x for x in (role.start, role.end) if x)
        if dates:
            p.add_run(f"\t{dates}")
        for b in role.bullets:
            _bullet(doc, b.text)

    if resume.projects:
        _heading(doc, "Projects")
        for pr in resume.projects:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.add_run(pr.name).bold = True
            label = {"portfolio": "Personal project", "learning": "Learning project", "professional": "Work project"}.get(pr.kind)
            if label and not any(w in pr.name.lower() for w in ("project", "portfolio", "personal", "course")):
                p.add_run(f" ({label})").italic = True
            if pr.description:
                doc.add_paragraph(pr.description)
            for b in pr.bullets:
                _bullet(doc, b.text)

    if resume.education:
        _heading(doc, "Education")
        for e in resume.education:
            p = doc.add_paragraph()
            p.add_run(e.degree or "").bold = True
            p.add_run((" | " if e.degree else "") + e.institution)
            dates = " – ".join(x for x in (e.start, e.end) if x)
            if dates:
                p.add_run(f" | {dates}")

    if resume.certifications:
        _heading(doc, "Certifications")
        for cert in resume.certifications:
            _bullet(doc, cert.name)

    doc.core_properties.author = resume.name
    doc.core_properties.title = title
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)
