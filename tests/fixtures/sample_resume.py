"""Builds a FICTIONAL sample resume used only for automated tests."""
from __future__ import annotations

from pathlib import Path

from docx import Document

SAMPLE = {
    "name": "Sample Candidate",
    "headline": "Data Engineer",
    "contact": "Lahore, Pakistan | sample.candidate@example.com | +92 300 0000000 | linkedin.com/in/sample-candidate",
    "summary": ("Data Engineer with 2+ years of experience building and maintaining production ETL pipelines "
                "using Python, SQL, PostgreSQL, AWS Lambda and AWS Step Functions."),
    "experience": [
        ("Data Engineer | Example PropTech Ltd | Lahore, Pakistan", "Jan 2023 – Present", [
            "Maintain production ETL pipelines in Python and SQL that ingest real-estate listing data from REST APIs and RETS feeds into Amazon RDS (PostgreSQL).",
            "Orchestrate ingestion and transformation workflows with AWS Lambda and AWS Step Functions across raw and staging layers.",
            "Implemented incremental loading and upsert logic to synchronise listings with target databases.",
            "Investigate pipeline failures and fix data quality issues, improving reliability of daily loads.",
            "Optimised slow PostgreSQL queries and indexes, reducing a nightly job runtime by 40%.",
        ]),
        ("Junior Data Analyst | Sample Analytics Co | Lahore, Pakistan", "Jun 2022 – Dec 2022", [
            "Built SSIS packages to load data from Microsoft SQL Server into reporting tables.",
            "Created Power BI dashboards for operations reporting.",
        ]),
        ("Data Intern | Demo Software House | Lahore, Pakistan", "Jan 2022 – May 2022", [
            "Wrote SQL queries for ad-hoc analysis.",
        ]),
    ],
    "projects": [
        ("Weather Data Lake (portfolio)", [
            "Built a portfolio pipeline loading weather API data to Amazon S3 and querying it with Amazon Athena.",
            "Scheduled tasks with Apache Airflow running in Docker.",
        ]),
    ],
    "skills": ["Languages: Python, SQL", "Databases: PostgreSQL, Microsoft SQL Server",
               "Cloud: AWS Lambda, AWS Step Functions, Amazon RDS", "BI: Power BI, Tableau",
               "Tools: SSIS, Git"],
    "education": [("BS Computer Science", "University of the Punjab, Lahore", "2018 – 2022")],
    "certifications": ["Databricks Lakehouse Fundamentals (Databricks, 2024)"],
}


def build_sample_resume(path: Path, data: dict = SAMPLE) -> Path:
    doc = Document()
    doc.add_paragraph(data["name"], style="Title")
    doc.add_paragraph(data["headline"])
    doc.add_paragraph(data["contact"])
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(data["summary"])
    doc.add_heading("Work Experience", level=1)
    for header, dates, bullets in data["experience"]:
        doc.add_paragraph(f"{header} | {dates}")
        for b in bullets:
            doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Projects", level=1)
    for name, bullets in data["projects"]:
        doc.add_paragraph(name)
        for b in bullets:
            doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Technical Skills", level=1)
    for s in data["skills"]:
        doc.add_paragraph(s)
    doc.add_heading("Education", level=1)
    for degree, inst, years in data["education"]:
        doc.add_paragraph(f"{degree} | {inst} | {years}")
    doc.add_heading("Certifications", level=1)
    for c in data["certifications"]:
        doc.add_paragraph(c)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
