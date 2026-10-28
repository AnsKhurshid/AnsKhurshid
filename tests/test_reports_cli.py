from pathlib import Path

import openpyxl
import yaml

from jobagent.cli import main
from jobagent.matching.analyzer import Analyzer
from jobagent.reporting.exports import generate_reports
from jobagent.storage.db import Database
from tests.conftest import JOBS, job_id_by_source_id


def test_reports_written_with_sample_size_disclosure(settings, taxonomy, profile, loaded_db):
    Analyzer(settings, loaded_db, taxonomy, profile).analyze()
    paths = generate_reports(settings, loaded_db)
    wb = openpyxl.load_workbook(paths["excel"])
    assert wb.sheetnames == ["All jobs", "Shortlisted", "Skill gaps"]
    assert wb["All jobs"].max_row == 15          # header + 14 unique jobs
    summary = Path(paths["summary_md"]).read_text()
    assert "based on 11 analysed jobs" in summary
    assert "not an exhaustive view" in summary
    csv = Path(paths["shortlist_csv"]).read_text(encoding="utf-8-sig")
    assert "Indus Data Systems" in csv and "Retail Insights" not in csv


def test_cli_end_to_end(settings, config_path, master_resume, capsys):
    cfg = ["--config", str(config_path)]
    assert main(cfg + ["init"]) == 0
    assert main(cfg + ["analyze"]) == 2                         # no profile yet -> refuses
    assert main(cfg + ["profile", "extract"]) == 0
    assert main(cfg + ["profile", "approve"]) == 2              # listed skills unresolved
    ppath = settings.path(settings.profile.candidate_profile_path)
    data = yaml.safe_load(ppath.read_text())
    for s in data["skills"]:
        if s["level"] == "listed":
            s["level"] = "studied"
    ppath.write_text(yaml.safe_dump(data, sort_keys=False))
    assert main(cfg + ["profile", "approve"]) == 0

    files = [str(JOBS / f) for f in ("jobs.json", "linkedin_saved_jsonld.html", "linkedin_saved_layout.html",
                                     "linkedin_urls.txt", "pasted_job.txt")]
    assert main(cfg + ["run", "--tailor", "--no-inbox", *files]) == 0
    out = capsys.readouterr().out
    assert "14 new" in out and "summary_md" in out and ": ready ->" in out

    db = Database(settings.path(settings.storage.db_path))
    stub = job_id_by_source_id(db, "3888888888")
    details = settings.root / "tmp_details.txt"
    try:
        details.write_text("Title: SQL Data Engineer\nCompany: Example Bank\nLocation: Karachi, Pakistan\n\n"
                           "Responsibilities:\n- Build SQL and Python ETL pipelines into PostgreSQL\n"
                           "Requirements:\n- 2+ years of experience\n- SQL, Python, PostgreSQL\n" + "More detail. " * 30)
        assert main(cfg + ["add-details", str(stub), str(details)]) == 0
    finally:
        details.unlink()
    assert main(cfg + ["analyze", "--job", str(stub)]) == 0
    assert db.latest_analysis(stub)["result"]["category"] in ("strong", "potential")

    assert main(cfg + ["set-status", str(stub), "applied", "--reason", "applied via company site"]) == 0
    assert db.job(stub)["status"] == "applied"
    assert main(cfg + ["show", str(stub)]) == 0
    assert main(cfg + ["jobs", "--status", "applied"]) == 0
    assert main(cfg + ["run", "--no-inbox", *files]) == 0       # rerun: no duplicates
    assert db.q1("SELECT COUNT(*) n FROM jobs")["n"] == 14
    assert db.job(stub)["status"] == "applied"

    assert main(cfg + ["purge", "--jobs", "--resumes", "--yes"]) == 0
    assert db.q1("SELECT COUNT(*) n FROM jobs")["n"] == 0
    assert master_resume.exists() and ppath.exists()
    assert not any(settings.path(settings.resume.output_dir).iterdir())
